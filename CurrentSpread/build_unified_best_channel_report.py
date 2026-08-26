from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\EM\CurrentSpread")
RESULTS = ROOT / "UnifiedBestChannelReanalysis_MaxOD_20260820"
SUMMARY_DIR = RESULTS / "summaries"
OUTPUT = RESULTS / "Unified_Best_Channel_Reanalysis_Summary_MaxOD_with_Population_Plots_and_Session_Counts_2026-08-20.docx"
PANEL_DIR = RESULTS / "figures" / "population_scatter_panels"

INK = "1F2937"
NAVY = "163A5F"
BLUE = "2E74B5"
MUTED = "64748B"
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GREEN = "215E45"
RED = "9B1C1C"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def set_run_font(run, size=11, bold=False, italic=False, color=INK, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_cell_text(cell, text, *, bold=False, color=INK, size=8.7, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)


def shade_paragraph(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def add_picture_with_alt(doc: Document, path: Path, width_inches: float, alt_text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(path), width=Inches(width_inches))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", path.stem)
    return paragraph


def add_custom_bullet_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_bullet(doc: Document, bullet_num_id: int, text: str, bold_lead: str | None = None):
    paragraph = doc.add_paragraph(style="Normal")
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(bullet_num_id))
    num_pr.extend([ilvl, num_id])
    p_pr.append(num_pr)
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = rgb(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)

    for name, size, before, after, color in (
        ("Heading 1", 16, 16, 8, BLUE),
        ("Heading 2", 13, 12, 6, BLUE),
        ("Heading 3", 12, 8, 4, NAVY),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_heading(doc: Document, text: str, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc: Document, text: str, bold_lead: str | None = None):
    paragraph = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_callout(doc: Document, label: str, text: str, fill=CALLOUT):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.line_spacing = 1.10
    shade_paragraph(paragraph, fill)
    label_run = paragraph.add_run(label + "  ")
    set_run_font(label_run, size=11, bold=True, color=NAVY)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=11, color=INK)
    return paragraph


def add_metadata(doc: Document, rows: list[tuple[str, str]]):
    for label, value in rows:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0
        left = paragraph.add_run(label + ": ")
        set_run_font(left, size=10.5, bold=True, color=INK)
        right = paragraph.add_run(value)
        set_run_font(right, size=10.5, color=INK)


def fnum(value: str, digits=3, na="N/A") -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return na
    if number != number:
        return na
    return f"{number:.{digits}f}"


def integer(value: str) -> str:
    try:
        return str(int(round(float(value))))
    except (TypeError, ValueError):
        return "N/A"


def find_run(rows, analysis: str, area: str) -> dict[str, str]:
    return next(row for row in rows if row["Analysis"] == analysis and row["Area"] == area)


def build_document():
    run_rows = read_csv(SUMMARY_DIR / "UnifiedRunSummary.csv")
    provenance = read_csv(SUMMARY_DIR / "InputProvenance.csv")[0]

    doc = Document()
    configure_styles(doc)
    bullet_id = add_custom_bullet_numbering(doc)
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run("UNIFIED BEST-CHANNEL REANALYSIS  |  MAX OD"), size=8.5, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.space_before = Pt(0)
    add_page_number(footer)

    title = doc.add_paragraph("Unified Best-Channel Population Reanalysis", style="Title")
    title.paragraph_format.keep_with_next = True
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_run_font(subtitle.add_run("Three harmonized reanalyses plus the original Quick comparison | Max OD only"), size=13, color=MUTED)
    add_metadata(doc, [
        ("Date", "August 20, 2026"),
        ("Audience", "Ari and the analysis team"),
        ("Shared input", provenance["InputFile"]),
        ("Input identity", f"251 rows; SHA-256 {provenance['SHA256']}"),
        ("Status", "Completed and validated"),
    ])

    add_callout(
        doc,
        "BOTTOM LINE",
        "None of the three main reanalyses produced a significant AI x OD interaction in either MT or FST, for either 2D or 3D neurons. The original Quick stimulation-channel comparison retains significant 2D interactions in both areas, but no 3D interaction. The new neural-source choices therefore did not improve the population result.",
        fill=PALE_BLUE,
    )

    add_heading(doc, "Main results", 1)
    add_body(doc, "Counts are shown after selected-source neural gating and after the common behavioral-fit requirements. The p-values are for AI:OD in the merged Dominant + NonDominant model.")
    headers = ["Analysis", "Area", "Source valid", "Neural 2D / 3D", "Model 2D / 3D", "p(AI x OD), 2D", "p(AI x OD), 3D"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, value in enumerate(headers):
        shade_cell(table.rows[0].cells[idx], LIGHT)
        set_cell_text(table.rows[0].cells[idx], value, bold=True, color=NAVY, size=8.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_header(table.rows[0])
    for row in run_rows:
        cells = table.add_row().cells
        values = [
            row["Analysis"], row["Area"], f"{integer(row['SourceValidRows'])}/{integer(row['SourceTableRows'])}",
            f"{integer(row['Neural2DRows'])} / {integer(row['Neural3DRows'])}",
            f"{integer(row['Model2DUnits'])} / {integer(row['Model3DUnits'])}",
            fnum(row["P_AIxOD_2D"]), fnum(row["P_AIxOD_3D"]),
        ]
        for idx, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cells[idx], value, size=8.2, align=align)
    set_table_geometry(table, [2300, 580, 1050, 1350, 1350, 1365, 1365])

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(0)
    set_run_font(note.add_run("Interpretation: "), size=9, bold=True, color=NAVY)
    set_run_font(note.add_run("among the three main reanalyses, all estimable interaction p-values are >= 0.063; the original comparator has significant 2D interactions in MT (p = 0.006) and FST (p = 0.002)."), size=9, color=MUTED)

    doc.add_page_break()
    add_heading(doc, "Shared specification applied to all four methods", 1)
    add_bullet(doc, bullet_id, "Same data table: all methods begin with the same 251 rows from unit_table_stim.mat. No workbook refresh or table join occurs in the population stage.", "Same data table:")
    add_bullet(doc, bullet_id, "Quick channel selection: the original comparator is fixed at StimElec; the two best-channel methods use maximum pooled Pearson r after one sample Z score across the complete 4-cue x 8-matched-coherence matrix; no SSE ranking.", "Quick channel selection:")
    add_bullet(doc, bullet_id, "Selected neural source: AI, Max OD, raw MonoL/MonoR p_AI, and Z3D-Z2D are all evaluated from the same selected task/channel for that analysis.", "Selected neural source:")
    add_bullet(doc, bullet_id, "OD convention: Max-response signed OD assigns dominant eye; only absolute OD enters regression weights, marker opacity, and summaries.", "OD convention:")
    add_bullet(doc, bullet_id, "Neural gate: raw MonoL p_AI < 0.05 and raw MonoR p_AI < 0.05. Z3D-Z2D < 0 defines 2D; > 0 defines 3D.", "Neural gate:")
    add_bullet(doc, bullet_id, "Behavioral gate: the NoStim and Stim sigmoid fits must both be good and finite for each plotted cue.", "Behavioral gate:")
    add_bullet(doc, bullet_id, "Anatomy: all AP values are retained; there is no experiment-specific AP cutoff.", "Anatomy:")
    add_bullet(doc, bullet_id, "Population model: NonDominant bias is sign-flipped, then Dominant + NonDominant points are fit with MergedEyeBias ~ AI + AI:OD.", "Population model:")
    add_bullet(doc, bullet_id, "Exclusions: no post hoc influence-based or session-outlier removal is applied.", "Exclusions:")

    add_heading(doc, "What differs among the four methods", 1)
    method_table = doc.add_table(rows=1, cols=3)
    method_table.style = "Table Grid"
    method_headers = ["Analysis", "Neural source/channel", "Selection boundary"]
    for idx, value in enumerate(method_headers):
        shade_cell(method_table.rows[0].cells[idx], LIGHT)
        set_cell_text(method_table.rows[0].cells[idx], value, bold=True, color=NAVY, size=9)
    set_repeat_header(method_table.rows[0])
    method_rows = [
        ("Original Quick - stimulation channel", "3DMotionQuick tuning at StimElec", "Fixed stimulation channel"),
        ("Stim-task NoStim", "3DMotionStim NoStim trials at StimElec", "Not applicable"),
        ("Best Quick - all channels", "3DMotionQuick channel with maximum pooled r", "All available contacts"),
        ("Best Quick - within 200 um", "3DMotionQuick channel with maximum pooled r", "+/-4 probe positions (+/-200 um)"),
    ]
    for values in method_rows:
        cells = method_table.add_row().cells
        for idx, value in enumerate(values):
            set_cell_text(cells[idx], value, size=9, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_table_geometry(method_table, [2300, 3900, 3160])

    add_heading(doc, "Channel-selection and 3D-count readout", 1)
    all_mt = find_run(run_rows, "Best Quick - all channels", "MT")
    all_fst = find_run(run_rows, "Best Quick - all channels", "FST")
    near_mt = find_run(run_rows, "Best Quick - within 200 um", "MT")
    near_fst = find_run(run_rows, "Best Quick - within 200 um", "FST")
    add_bullet(doc, bullet_id, f"All-channel Quick selection: median r = {fnum(all_mt['MedianSelectionR'])} in MT and {fnum(all_fst['MedianSelectionR'])} in FST; median absolute distance = {integer(all_mt['MedianAbsDistanceMicrometers'])} and {integer(all_fst['MedianAbsDistanceMicrometers'])} um.", "All-channel Quick selection:")
    add_bullet(doc, bullet_id, f"Within-200-um selection: median r = {fnum(near_mt['MedianSelectionR'])} in MT and {fnum(near_fst['MedianSelectionR'])} in FST; median absolute distance = {integer(near_mt['MedianAbsDistanceMicrometers'])} and {integer(near_fst['MedianAbsDistanceMicrometers'])} um.", "Within-200-um selection:")
    add_bullet(doc, bullet_id, "The 200-um boundary changed 13/91 MT selections and 28/160 FST selections overall. Twelve of the MT changes and all 28 FST changes remained source-valid for the unified population preparation.", "Selection changes:")
    add_bullet(doc, bullet_id, "Restricting distance increased FST modeled 3D units from 19 to 22, but p(AI x OD) changed only from 0.673 to 0.643 and remained nonsignificant. MT retained 4 modeled 3D units in both Quick analyses.", "3D units:")

    add_heading(doc, "Interpretation", 1)
    add_callout(doc, "CONCLUSION", "The original Quick stimulation-channel method retains significant 2D interactions in MT and FST, but not a 3D interaction. Under the harmonized specification, none of the three main alternative neural-source choices recovers or strengthens that population effect.", fill=CALLOUT)
    add_bullet(doc, bullet_id, "The closest interaction was FST 2D with the all-channel Quick method (p = 0.063). It is still above 0.05 and is not the intended 3D result.", "Closest result:")
    add_bullet(doc, bullet_id, "Stim-task NoStim recomputation leaves only 1 modeled MT 3D unit and 7 modeled FST 3D units, making it the weakest source for a 3D population test.", "Stim-task 3D cohort:")
    add_bullet(doc, bullet_id, "The within-200-um rule modestly improves the FST 3D count but does not improve the inference. This argues against distance alone as the explanation.", "Distance restriction:")
    add_bullet(doc, bullet_id, "The next useful inspection is targeted session QC for source failures, surprising selected-channel class changes, and leverage—not another post hoc deletion pass aimed at significance.", "Next diagnostic:")

    add_heading(doc, "Explicit source exceptions", 1)
    add_body(doc, "Stim-task NoStim: row 19 has undefined Max OD; rows 21, 34, 48, 79, and 184 have partial AI/OD support. These rows remain in the source audit but are not source-valid for the unified population input.")
    add_body(doc, "Quick methods: channel correlation succeeded for all 251 rows in both best-channel analyses. Row 102 cannot recompute selected-channel p_AI/Z3D-Z2D because C:\\Jim\\StimData\\20240524.mat is missing; it is excluded identically from the original comparator and both best-channel population inputs.")

    add_heading(doc, "Clean result package", 1)
    output_rows = [
        ("UnifiedRunSummary.csv", "One row per analysis and area; fastest comparison of counts and model results."),
        ("UnifiedAIxODSummary.csv", "Complete merged-eye AI, AI:OD, and R2 output by analysis, area, and unit type."),
        ("UnifiedCriteria.csv", "Machine-readable shared analysis specification."),
        ("UnifiedSessionAudit.csv", "1,004 method-session rows with selections and neural-gate validity."),
    ]
    package_table = doc.add_table(rows=1, cols=2)
    package_table.style = "Table Grid"
    for idx, value in enumerate(("File", "Purpose")):
        shade_cell(package_table.rows[0].cells[idx], LIGHT)
        set_cell_text(package_table.rows[0].cells[idx], value, bold=True, color=NAVY, size=9)
    set_repeat_header(package_table.rows[0])
    for filename, purpose in output_rows:
        cells = package_table.add_row().cells
        set_cell_text(cells[0], filename, bold=True, color=NAVY, size=7.9)
        set_cell_text(cells[1], purpose, size=8.8)
    set_table_geometry(package_table, [3900, 5460])

    landscape = doc.add_section(WD_SECTION_START.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width = Inches(11)
    landscape.page_height = Inches(8.5)
    landscape.top_margin = Inches(0.4)
    landscape.right_margin = Inches(0.4)
    landscape.bottom_margin = Inches(0.4)
    landscape.left_margin = Inches(0.4)
    landscape.header_distance = Inches(0.20)
    landscape.footer_distance = Inches(0.20)

    figure_methods = [
        (
            "Original comparison - Quick tuning at stimulation channel",
            PANEL_DIR / "00_OriginalQuick_StimChannel_Population_FourPanel.png",
            "Four population scatter plots for the original Quick stimulation-channel comparison. Top row: MT 2D and MT 3D. Bottom row: FST 2D and FST 3D. Each panel lists included-session counts for all four cues.",
        ),
        (
            "Stim-task NoStim tuning",
            PANEL_DIR / "01_StimTask_NoStim_Population_FourPanel.png",
            "Four population scatter plots for Stim-task NoStim tuning. Top row: MT 2D and MT 3D. Bottom row: FST 2D and FST 3D. Each panel lists included-session counts for all four cues.",
        ),
        (
            "Best Quick channel - all channels",
            PANEL_DIR / "02_BestQuick_AllChannels_Population_FourPanel.png",
            "Four population scatter plots for the best Quick channel across all channels. Top row: MT 2D and MT 3D. Bottom row: FST 2D and FST 3D. Each panel lists included-session counts for all four cues.",
        ),
        (
            "Best Quick channel - within 200 um",
            PANEL_DIR / "03_BestQuick_Within200um_Population_FourPanel.png",
            "Four population scatter plots for the best Quick channel within 200 micrometers. Top row: MT 2D and MT 3D. Bottom row: FST 2D and FST 3D. Each panel lists included-session counts for all four cues.",
        ),
    ]
    for index, (label, image_path, alt_text) in enumerate(figure_methods):
        if index:
            doc.add_page_break()
        heading = add_heading(doc, "Population scatter plots | " + label, 2)
        heading.paragraph_format.space_before = Pt(0)
        heading.paragraph_format.space_after = Pt(2)
        order = doc.add_paragraph()
        order.paragraph_format.space_before = Pt(0)
        order.paragraph_format.space_after = Pt(2)
        set_run_font(order.add_run("Panel order: MT 2D (top left), MT 3D (top right), FST 2D (bottom left), FST 3D (bottom right). Count boxes report included sessions by cue."), size=8.5, color=MUTED)
        add_picture_with_alt(doc, image_path, 8.65, alt_text)

    doc.core_properties.title = "Unified Best-Channel Population Reanalysis - Max OD"
    doc.core_properties.subject = "Shared-criteria comparison of the original Quick stimulation-channel method, Stim-task NoStim, and two Best Quick channel analyses"
    doc.core_properties.author = "Analysis team"
    doc.core_properties.keywords = "best channel, population analysis, Max OD, Pearson correlation, 2D, 3D"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
