from __future__ import annotations

from pathlib import Path
import math

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_ROOT = Path(r"C:\EM\CurrentSpread")
BUILD_DIR = Path(
    r"C:\Users\zzhu329\.codex\visualizations\2026\08\19\01a01a99-5cf1-7523-b616-a0f87614ae34"
) / "best_channel_report_work"
OUTPUT_DOCX = OUTPUT_ROOT / "Best_Channel_Analysis_Summary_2026-08-20.docx"
HISTOGRAM_PNG = BUILD_DIR / "drift_distance_histograms_by_monkey_area.png"

SESSION_SUMMARY = Path(
    r"C:\EM\StimTuningAnalysis\QuickVsStimNoStimCorrelation\QuickStimNoStim_SessionSummary.csv"
)
STIM_MANIFEST = Path(r"C:\EM\StimTuningAnalysis\StimTuningSessionManifest.csv")
WITHIN200_AIXOD = Path(
    r"C:\EM\StimTuningAnalysis\BestCorrelationChannelPopulation_Within200um_ODComparison\AIxODSummary_Max_vs_LSQ_MT_FST.csv"
)
WITHIN200_RUN = Path(
    r"C:\EM\StimTuningAnalysis\BestCorrelationChannelPopulation_Within200um_ODComparison\RunSummary_Max_vs_LSQ_MT_FST.csv"
)
INFLUENCE_SUMMARY = Path(
    r"C:\EM\StimTuningAnalysis\BestQuickChannelPopulation\OutlierAnalysis_SelectedChannelCriteria\InfluenceModelSummary.csv"
)
QUICK_CLEAN_MANIFEST = Path(
    r"C:\EM\QuickTuningCleaningAnalysis\3DQuick\3DQuick_CleaningSessionManifest.csv"
)


BLUE = "2E74B5"
DARK_BLUE = "17365D"
INK = "243447"
MUTED = "657386"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
GRID = "CBD5E1"
WHITE = "FFFFFF"
RED = "9B1C1C"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, size=11, bold=False, italic=False, color=INK, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = int(sum(widths_dxa))
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width)))
        grid.append(col)

    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(widths_dxa[idx])))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text, bold=False, color=INK, size=9, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)


def set_paragraph_shading(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_left_border(paragraph, color=BLUE, size=18, space=8):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    set_run_font(run)
    return p


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = rgb(DARK_BLUE)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(13)
    subtitle.font.italic = False
    subtitle.font.color.rgb = rgb(MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(12)

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    list_style = styles["List Bullet"]
    list_style.font.name = "Calibri"
    list_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    list_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    list_style.font.size = Pt(11)
    list_style.paragraph_format.left_indent = Inches(0.5)
    list_style.paragraph_format.first_line_indent = Inches(-0.25)
    list_style.paragraph_format.space_after = Pt(5)


def load_drift_data():
    summary = pd.read_csv(SESSION_SUMMARY)
    manifest = pd.read_csv(STIM_MANIFEST)
    merged = summary.merge(
        manifest[["UnitTableRow", "TInfoFile"]], on="UnitTableRow", how="left", validate="one_to_one"
    )
    merged["Area"] = merged["TInfoFile"].str.extract(r"_(MT|FST)_", expand=False)
    if merged["Area"].isna().any():
        missing = merged.loc[merged["Area"].isna(), "UnitTableRow"].tolist()
        raise RuntimeError(f"Could not infer area for rows: {missing}")
    merged["DistanceUm"] = pd.to_numeric(merged["BestDistanceToStimMicrometers"])
    merged["AbsDistanceUm"] = merged["DistanceUm"].abs()
    return merged


def drift_summary_table(data):
    rows = []
    group_order = [("Jim", "MT"), ("Jim", "FST"), ("Clay", "MT"), ("Clay", "FST")]
    for monkey, area in group_order:
        group = data[(data["Monkey"] == monkey) & (data["Area"] == area)]
        n = len(group)
        rows.append(
            {
                "Group": f"{monkey} - {area}",
                "N": n,
                "MedianAbs": int(group["AbsDistanceUm"].median()),
                "Within100": f"{(group['AbsDistanceUm'] <= 100).sum()} ({100 * (group['AbsDistanceUm'] <= 100).mean():.1f}%)",
                "Within200": f"{(group['AbsDistanceUm'] <= 200).sum()} ({100 * (group['AbsDistanceUm'] <= 200).mean():.1f}%)",
                "Range": f"{int(group['DistanceUm'].min())} to {int(group['DistanceUm'].max())}",
            }
        )
    n = len(data)
    rows.append(
        {
            "Group": "All sessions",
            "N": n,
            "MedianAbs": int(data["AbsDistanceUm"].median()),
            "Within100": f"{(data['AbsDistanceUm'] <= 100).sum()} ({100 * (data['AbsDistanceUm'] <= 100).mean():.1f}%)",
            "Within200": f"{(data['AbsDistanceUm'] <= 200).sum()} ({100 * (data['AbsDistanceUm'] <= 200).mean():.1f}%)",
            "Range": f"{int(data['DistanceUm'].min())} to {int(data['DistanceUm'].max())}",
        }
    )
    return rows


def create_histogram(data):
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    colors = {
        ("Jim", "MT"): "2E74B5",
        ("Jim", "FST"): "1F9E89",
        ("Clay", "MT"): "E58C3A",
        ("Clay", "FST"): "C55454",
    }
    order = [("Jim", "MT"), ("Jim", "FST"), ("Clay", "MT"), ("Clay", "FST")]
    distance_values = list(range(-750, 751, 50))
    histograms = {}
    max_count = 0
    for key in order:
        values = data.loc[(data["Monkey"] == key[0]) & (data["Area"] == key[1]), "DistanceUm"]
        value_counts = values.value_counts()
        counts = [int(value_counts.get(distance, 0)) for distance in distance_values]
        histograms[key] = (values, counts)
        max_count = max(max_count, max(counts))

    image = Image.new("RGB", (2200, 1370), "white")
    draw = ImageDraw.Draw(image)

    def font(size, bold=False):
        filename = "arialbd.ttf" if bold else "arial.ttf"
        path = Path(r"C:\Windows\Fonts") / filename
        return ImageFont.truetype(str(path), size=size)

    title_font = font(38, bold=True)
    panel_title_font = font(27, bold=True)
    body_font = font(21)
    small_font = font(18)
    tick_font = font(17)
    axis_font = font(20, bold=True)

    title = "Unconstrained best-channel displacement by monkey and cortical area"
    title_box = draw.textbbox((0, 0), title, font=title_font)
    draw.text(((2200 - (title_box[2] - title_box[0])) / 2, 28), title, font=title_font, fill="#17365D")

    panels = [(70, 125, 1060, 625), (1140, 125, 2130, 625), (70, 705, 1060, 1205), (1140, 705, 2130, 1205)]
    y_limit = int(math.ceil((max_count + 2) / 5) * 5)
    x_ticks = [-750, -500, -250, 0, 250, 500, 750]

    for key, panel in zip(order, panels):
        x0, y0, x1, y1 = panel
        plot_left, plot_top = x0 + 85, y0 + 74
        plot_right, plot_bottom = x1 - 25, y1 - 72
        plot_width = plot_right - plot_left
        plot_height = plot_bottom - plot_top
        values, counts = histograms[key]

        draw.text((x0 + 6, y0 + 3), f"{key[0]} - {key[1]}", font=panel_title_font, fill="#17365D")
        stats = f"N = {len(values)}   median |d| = {values.abs().median():.0f} µm"
        stats_box = draw.textbbox((0, 0), stats, font=small_font)
        draw.text((x1 - 8 - (stats_box[2] - stats_box[0]), y0 + 9), stats, font=small_font, fill="#657386")

        for y_tick in range(0, y_limit + 1, 5):
            y = plot_bottom - (y_tick / y_limit) * plot_height
            draw.line((plot_left, y, plot_right, y), fill="#D9E1E8", width=2)
            label = str(y_tick)
            box = draw.textbbox((0, 0), label, font=tick_font)
            draw.text((plot_left - 12 - (box[2] - box[0]), y - (box[3] - box[1]) / 2), label, font=tick_font, fill="#657386")

        def x_position(distance):
            return plot_left + ((distance + 775) / 1550) * plot_width

        bar_slot = plot_width / len(distance_values)
        for distance, count in zip(distance_values, counts):
            left = x_position(distance - 25) + 1
            right = x_position(distance + 25) - 1
            top = plot_bottom - (count / y_limit) * plot_height
            if count > 0:
                draw.rectangle((left, top, right, plot_bottom), fill=f"#{colors[key]}", outline="white", width=2)

        for threshold in (-200, 200):
            x = x_position(threshold)
            dash = 12
            y = plot_top
            while y < plot_bottom:
                draw.line((x, y, x, min(y + dash, plot_bottom)), fill="#7A5A00", width=3)
                y += dash * 2
        zero_x = x_position(0)
        draw.line((zero_x, plot_top, zero_x, plot_bottom), fill="#263238", width=3)

        draw.line((plot_left, plot_top, plot_left, plot_bottom), fill="#94A3B8", width=3)
        draw.line((plot_left, plot_bottom, plot_right, plot_bottom), fill="#94A3B8", width=3)
        for tick in x_ticks:
            x = x_position(tick)
            draw.line((x, plot_bottom, x, plot_bottom + 7), fill="#94A3B8", width=2)
            label = str(tick)
            box = draw.textbbox((0, 0), label, font=tick_font)
            draw.text((x - (box[2] - box[0]) / 2, plot_bottom + 11), label, font=tick_font, fill="#657386")

        draw.text((x0 + 5, plot_top - 30), "Sessions", font=axis_font, fill="#4B5563")
        if key in (("Clay", "MT"), ("Clay", "FST")):
            x_label = "Signed displacement (µm)"
            box = draw.textbbox((0, 0), x_label, font=axis_font)
            draw.text(((plot_left + plot_right - (box[2] - box[0])) / 2, y1 - 33), x_label, font=axis_font, fill="#4B5563")

    foot = "Dashed lines mark ±200 µm. One contact position corresponds to 50 µm."
    foot_box = draw.textbbox((0, 0), foot, font=body_font)
    draw.text(((2200 - (foot_box[2] - foot_box[0])) / 2, 1318), foot, font=body_font, fill="#657386")
    image.save(HISTOGRAM_PNG, format="PNG", optimize=True)


def format_p(value):
    value = float(value)
    if value < 0.001:
        return "<.001"
    return f"{value:.3f}".replace("0.", ".")


def build_document(data):
    aixod = pd.read_csv(WITHIN200_AIXOD)
    run_summary = pd.read_csv(WITHIN200_RUN)
    influence = pd.read_csv(INFLUENCE_SUMMARY)
    quick_clean = pd.read_csv(QUICK_CLEAN_MANIFEST)

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    r = header.add_run("BEST-CHANNEL ANALYSIS  |  INTERNAL RESEARCH SUMMARY")
    set_run_font(r, size=8.5, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph(style="Title")
    title.add_run("Best-Channel and Drift Analysis")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("What was tested, what changed, and why the population result did not improve")

    metadata = doc.add_paragraph()
    metadata.paragraph_format.space_after = Pt(12)
    for label, value in (
        ("Prepared for", "Discussion with Ari"),
        ("Date", "August 20, 2026"),
        ("Dataset", "251 paired 3DMotionQuick / 3DMotionStim sessions; Jim and Clay; MT and FST"),
    ):
        run = metadata.add_run(f"{label}: ")
        set_run_font(run, size=9.5, bold=True, color=DARK_BLUE)
        run = metadata.add_run(value + "\n")
        set_run_font(run, size=9.5, color=MUTED)

    callout = doc.add_paragraph()
    callout.paragraph_format.left_indent = Inches(0.16)
    callout.paragraph_format.right_indent = Inches(0.12)
    callout.paragraph_format.space_before = Pt(2)
    callout.paragraph_format.space_after = Pt(12)
    callout.paragraph_format.line_spacing = 1.12
    set_paragraph_shading(callout, LIGHT_BLUE)
    set_paragraph_left_border(callout)
    run = callout.add_run("Bottom line. ")
    set_run_font(run, size=11, bold=True, color=DARK_BLUE)
    run = callout.add_run(
        "Across the completed sensitivity analyses, replacing the stimulation-channel Quick tuning "
        "with a best-matching channel did not recover or strengthen the intended 3D neural-behavior "
        "population relationship. Method refinements changed selected channels, OD assignments, and "
        "cohort size, but not the central conclusion."
    )
    set_run_font(run, size=11, color=INK)

    add_heading(doc, "Completed analyses and outcomes", 1)
    attempts = [
        (
            "Stim-task NoStim tuning",
            "Replaced Quick-derived AI/OD with tuning measured during non-electrical-stimulation trials on the stimulation channel.",
            "Did not improve the target population relationship.",
        ),
        (
            "Best Quick channel",
            "Matched each Stim NoStim reference to all Quick channels using pooled whole-channel Z-scored tuning.",
            "Median best r = .730; the stimulation channel won 42/251 sessions. No population improvement.",
        ),
        (
            "Pearson r versus SSE",
            "Ranked channels separately by maximum correlation and minimum squared error.",
            "Identical winner in 251/251 sessions; after Z-scoring, the metrics are monotonic equivalents.",
        ),
        (
            "Selected-channel neural criteria",
            "Recomputed p_AI and Z3D-Z2D from the selected channel instead of retaining stimulation-channel labels.",
            "Unconstrained FST qualified 3D count fell from 31 to 21; the intended 3D effect remained weak.",
        ),
        (
            "Max OD versus LSQ OD",
            "Defined the dominant eye by peak response or by the lower Combined-to-Mono residual; modeled |OD| only.",
            "Eye assignments changed, but neither OD definition recovered the 3D population result.",
        ),
        (
            "Full FST cohort",
            "Removed the experiment-specific AP <= 26 restriction and reran the selected-channel population.",
            "Restored valid sessions but did not restore the target population relationship.",
        ),
        (
            "Influence / outlier sessions",
            "Computed leave-one-session-out influence and inspected the strongest flattening candidates.",
            "Removing the top 3D influence case changed p(AI×OD) from .673 to .799; no rescue of the 3D result.",
        ),
        (
            "Best channel within 200 µm",
            "Restricted correlation selection to ±4 contacts and reran MT/FST under both OD definitions.",
            "Removed extreme MT points, but all eight 2D/3D AI×OD tests remained nonsignificant.",
        ),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Analysis tried", "What changed", "Result"]
    for idx, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], text, bold=True, color=WHITE, size=9.2)
        set_cell_shading(table.rows[0].cells[idx], DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for attempt, change, result in attempts:
        cells = table.add_row().cells
        set_cell_text(cells[0], attempt, bold=True, color=DARK_BLUE, size=8.7)
        set_cell_text(cells[1], change, size=8.6)
        set_cell_text(cells[2], result, size=8.6)
    set_table_geometry(table, [2050, 3580, 3730])

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(8)
    r = note.add_run("Interpretive note: ")
    set_run_font(r, size=8.5, bold=True, color=MUTED)
    r = note.add_run(
        "The influence-based exclusion strengthened a 2D diagnostic interaction, but it was post hoc and did not improve the primary 3D result."
    )
    set_run_font(r, size=8.5, italic=True, color=MUTED)

    doc.add_page_break()
    add_heading(doc, "Final 200-µm sensitivity analysis", 1)
    add_body(
        doc,
        "This was the most constrained completed version: correlation-only channel selection within ±200 µm, selected-channel neural criteria, no AP cutoff, and unsigned OD in the models.",
    )
    endpoint_table = doc.add_table(rows=1, cols=6)
    endpoint_table.style = "Table Grid"
    endpoint_headers = ["OD", "Area", "Neural 2D / 3D", "Model 2D / 3D", "p(AI×OD), 2D", "p(AI×OD), 3D"]
    for idx, text in enumerate(endpoint_headers):
        set_cell_text(endpoint_table.rows[0].cells[idx], text, bold=True, color=WHITE, size=8.3, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(endpoint_table.rows[0].cells[idx], BLUE)
    set_repeat_table_header(endpoint_table.rows[0])
    for od_method, area in (("Max", "MT"), ("LSQ", "MT"), ("Max", "FST"), ("LSQ", "FST")):
        rs = run_summary[(run_summary["ODMethod"] == od_method) & (run_summary["Area"] == area)].iloc[0]
        r2d = aixod[(aixod["ODMethod"] == od_method) & (aixod["RequestedArea"] == area) & (aixod["UnitType"] == "2D")].iloc[0]
        r3d = aixod[(aixod["ODMethod"] == od_method) & (aixod["RequestedArea"] == area) & (aixod["UnitType"] == "3D")].iloc[0]
        values = [
            od_method,
            area,
            f"{int(rs.Neural2DCount)} / {int(rs.Neural3DCount)}",
            f"{int(rs.Model2DUnits)} / {int(rs.Model3DUnits)}",
            format_p(r2d.P_AIxOD),
            format_p(r3d.P_AIxOD),
        ]
        cells = endpoint_table.add_row().cells
        for idx, value in enumerate(values):
            set_cell_text(cells[idx], value, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        if area == "FST":
            for cell in cells:
                set_cell_shading(cell, LIGHT_BLUE)
    set_table_geometry(endpoint_table, [900, 700, 1800, 1700, 2130, 2130])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(
        "All p(AI×OD) values are nonsignificant. MT 3D remains especially underpowered (4 model units); FST has 22 modeled 3D units but still shows no interaction under either OD definition."
    )
    set_run_font(r, size=9.2, bold=True, color=RED)

    add_heading(doc, "What changed under the distance restriction", 2)
    add_bullet(
        doc,
        "The 200-µm rule changed 13/91 MT selections and 28/160 FST selections; the median selected distance became 50 µm in both areas.",
    )
    add_bullet(
        doc,
        "The constrained FST cohort contained 24 neural-qualified 3D sessions and 22 sessions with usable behavioral fits, compared with 21 neural-qualified 3D sessions in the earlier unconstrained selected-channel run.",
    )
    add_bullet(
        doc,
        "Removing distant winners eliminated two visibly extreme MT points, but it did not change the statistical conclusion in either area or under either OD definition.",
    )

    doc.add_page_break()
    add_heading(doc, "Inferred drift-distance distributions", 1)
    add_body(
        doc,
        "The histograms below use the unconstrained best-correlation channel from all 251 sessions. Distance is signed in physical probe order; zero means the best Quick channel was the stimulation channel. The measure is an inferred matched-channel displacement, not a direct measurement of mechanical probe motion.",
    )
    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.keep_with_next = True
    run = picture_paragraph.add_run()
    inline = run.add_picture(str(HISTOGRAM_PNG), width=Inches(6.55))
    inline._inline.docPr.set(
        "descr",
        "Four histograms of signed best-channel displacement, separated into Jim MT, Jim FST, Clay MT, and Clay FST sessions. Dashed lines mark plus or minus 200 micrometers.",
    )

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.keep_with_next = True
    r = caption.add_run(
        "Figure 1. Unconstrained best-channel displacement. Bin width = 50 µm (one contact); dashed lines = ±200 µm sensitivity threshold."
    )
    set_run_font(r, size=8.5, italic=True, color=MUTED)

    drift_rows = drift_summary_table(data)
    drift_table = doc.add_table(rows=1, cols=6)
    drift_table.style = "Table Grid"
    drift_headers = ["Group", "N", "Median |d| (µm)", "Within 100 µm", "Within 200 µm", "Signed range (µm)"]
    for idx, text in enumerate(drift_headers):
        set_cell_text(drift_table.rows[0].cells[idx], text, bold=True, color=WHITE, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(drift_table.rows[0].cells[idx], DARK_BLUE)
    set_repeat_table_header(drift_table.rows[0])
    for row_idx, row in enumerate(drift_rows):
        cells = drift_table.add_row().cells
        values = [row["Group"], row["N"], row["MedianAbs"], row["Within100"], row["Within200"], row["Range"]]
        for idx, value in enumerate(values):
            set_cell_text(cells[idx], value, size=8.2, bold=(row["Group"] == "All sessions"), align=WD_ALIGN_PARAGRAPH.CENTER)
        if row_idx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, LIGHT_GRAY)
        if row["Group"] == "All sessions":
            for cell in cells:
                set_cell_shading(cell, LIGHT_BLUE)
    set_table_geometry(drift_table, [1450, 620, 1540, 1800, 1800, 2150])

    doc.add_page_break()
    add_heading(doc, "Current interpretation", 1)
    add_bullet(
        doc,
        "Best-channel selection optimized task-to-task tuning similarity, not neural-behavior prediction. Better tuning agreement therefore did not guarantee a stronger population relationship.",
    )
    add_bullet(
        doc,
        "Replacing the stimulation channel also changed the biological unit being classified. The reduction and turnover in the 3D cohort are consequences of that redefinition, not merely loss of statistical power from a fixed cohort.",
    )
    add_bullet(
        doc,
        "Large inferred displacements should be interpreted cautiously: a distant channel may win because of shared tuning or noise, not because the probe physically drifted by that amount.",
    )
    add_bullet(
        doc,
        "The completed results currently support treating best-channel correction as a negative sensitivity analysis rather than as the primary population analysis.",
    )

    add_heading(doc, "Work not yet counted as a completed population test", 1)
    success_count = int((quick_clean["Status"] == "Success").sum())
    failed_count = int((quick_clean["Status"] != "Success").sum())
    add_body(
        doc,
        f"Quick-tuning outlier cleaning has generated cleaned 3D Quick results for {success_count}/251 sessions ({failed_count} unresolved failure). The current best-channel population runner still points to unit_table_stim.mat, so a fully documented population rerun from the cleaned Quick table is not yet represented above.",
    )
    add_body(
        doc,
        "A proposed next QC rule - requiring significant adjacent channels to have the same Combined-cue AI sign as the stimulation channel - has been discussed but not yet tested. It should remain separate from the completed negative analyses.",
    )

    add_heading(doc, "Analysis snapshot and source artifacts", 2)
    source_note = (
        "Snapshot: August 20, 2026. Drift distribution: unconstrained QuickStimNoStim session summary joined to the stimulation manifest. "
        "Population endpoint: 200-µm Max-versus-LSQ MT/FST aggregate summaries. Influence result: selected-channel session-influence summary. "
        "Quick-cleaning status: 3D Quick cleaning manifest."
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(source_note)
    set_run_font(r, size=8.2, color=MUTED)

    core = doc.core_properties
    core.title = "Best-Channel and Drift Analysis Summary"
    core.subject = "Summary of channel-drift analyses and population-level outcomes"
    core.author = "EM-3DMotion analysis team"
    core.keywords = "best channel, drift, MT, FST, population analysis"

    doc.save(OUTPUT_DOCX)


def main():
    required = [
        SESSION_SUMMARY,
        STIM_MANIFEST,
        WITHIN200_AIXOD,
        WITHIN200_RUN,
        INFLUENCE_SUMMARY,
        QUICK_CLEAN_MANIFEST,
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing required source files:\n" + "\n".join(missing))
    data = load_drift_data()
    if len(data) != 251:
        raise RuntimeError(f"Expected 251 drift rows, found {len(data)}")
    create_histogram(data)
    build_document(data)
    print(f"Created {OUTPUT_DOCX}")
    print(f"Created {HISTOGRAM_PNG}")


if __name__ == "__main__":
    main()
